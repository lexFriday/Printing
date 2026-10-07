from pathlib import Path
from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.section import WD_SECTION

root = Path('NJL_operations_visual_review')
out = Path('NJL_operations_visual_review.docx')
doc = Document()
sec = doc.sections[0]
sec.top_margin = Inches(.55); sec.bottom_margin = Inches(.55); sec.left_margin = Inches(.65); sec.right_margin = Inches(.65)
title = doc.add_heading('NJL Operations visual review', 0)
title.alignment = WD_ALIGN_PARAGRAPH.CENTER
p = doc.add_paragraph('Annotated prototype review for workflow, tracking, and cognitive load. ')
p.add_run('Ponytail principle: ').bold = True; p.add_run('keep one next action and avoid duplicate workflow state.')
items = [
    ('Calendar', 'calendar.png', 'Calendar is Job Order-linked', 'Selected date shows status, customer, target completion, and Open Job Order.', 'Add blocker/next-action text beside the date card.'),
    ('Inventory', 'inventory.png', 'Balance list -> operational view', 'Reorder warnings, linked jobs, and movement context.', 'Add one clear Open material action.'),
    ('Purchasing', 'purchasing.png', 'Requests now show progress', 'Supplier, remaining quantity, receipt progress, and job link.', 'Make the status action dominant.'),
    ('Production / BOM', 'production-bom.png', 'BOM templates are editable', 'Versioned sample recipe and demo-data warning.', 'Move users into the Job Order context before release.'),
]
for i, (heading, image, changed, added, improve) in enumerate(items):
    if i: doc.add_page_break()
    doc.add_heading(heading, 1)
    img = root / image
    if img.exists(): doc.add_picture(str(img), width=Inches(6.2))
    for label, value in [('What changed', changed), ('What was added', added), ('Smallest improvement', improve)]:
        p = doc.add_paragraph(); p.paragraph_format.space_after = Pt(3)
        r = p.add_run(label + ': '); r.bold = True; r.font.color.rgb = RGBColor(120,53,15); p.add_run(value)
doc.add_page_break(); doc.add_heading('Review conclusion', 1)
doc.add_paragraph('The four screens now read as connected operational documents: Calendar points to the Job Order, Production explains readiness, Inventory explains stock truth, and Purchasing explains shortage resolution. The next UI improvement is to keep the blocker and next action visible at the same point where the user decides what to do.')
doc.save(out)
print(out.resolve())
