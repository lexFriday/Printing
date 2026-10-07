from datetime import date, datetime
from pathlib import Path
import openpyxl

ROOT = Path(__file__).resolve().parents[1]
BATCH = "ACI-2026-legacy-v1"

def sql(value):
    if value is None or value == "":
        return "NULL"
    if isinstance(value, datetime):
        value = value.date().isoformat()
    elif isinstance(value, date):
        value = value.isoformat()
    if isinstance(value, (int, float)):
        return str(round(float(value), 2))
    return "'" + str(value).strip().replace("'", "''") + "'"

def parse_date(value, fallback=None):
    if isinstance(value, datetime):
        return value.date().isoformat()
    if isinstance(value, date):
        return value.isoformat()
    if isinstance(value, str):
        for fmt in ("%b. %d, %Y", "%B %d, %Y", "%m/%d/%Y", "%m/%d/%y"):
            try:
                return datetime.strptime(value.strip(), fmt).date().isoformat()
            except ValueError:
                pass
    return fallback

def expense_type(name, description, category=None):
    value = f"{category or ''} {name or ''} {description or ''}".lower()
    if any(word in value for word in ("material", "syringe", "plywood", "hardware", "gas")):
        return "materials"
    if any(word in value for word in ("transport", "taxi", "lalamove", "transportify", "parking", "van rental")):
        return "delivery"
    if any(word in value for word in ("labor", "salary", "allowance", "cash advance", "13th month")):
        return "labor"
    if "outsourc" in value:
        return "outsourcing"
    return "miscellaneous"

columns = "entry_type, entry_date, reference_no, document_no, counterparty_name, description, expense_type, gross_amount, vat_amount, ewt_amount, tin, address, source_workbook, source_sheet, source_row, import_batch, notes"
rows = []
expense_rows = []
count = 0

sales_file = Path(r"C:\Users\skhart\Downloads\ACI CASH SALES 2026.xlsx")
sales_workbook = openpyxl.load_workbook(sales_file, data_only=True, read_only=True)
for sheet in sales_workbook.worksheets:
    for row_number, row in enumerate(sheet.iter_rows(values_only=True), 1):
        if row_number == 1 or len(row) < 4 or not row[2] or str(row[2]).strip().upper() == "TOTAL" or not isinstance(row[3], (int, float)):
            continue
        entry_date = parse_date(row[1])
        values = ["revenue", entry_date, row[0], row[4], str(row[2]).strip(), "Cash sale / collection", None, row[3], 0, row[6] if isinstance(row[6], (int, float)) else 0, row[7], row[8], sales_file.name, sheet.title, row_number, BATCH, "Imported from ACI CASH SALES 2026.xlsx"]
        rows.append(f"({', '.join(sql(value) for value in values)})")
        count += 1

expense_file = Path(r"C:\Users\skhart\Downloads\ACI CASH DISBURSEMENT 2026.xlsx")
expense_workbook = openpyxl.load_workbook(expense_file, data_only=True, read_only=True)
for sheet in expense_workbook.worksheets:
    current_date = None
    current_category = None
    for row_number, row in enumerate(sheet.iter_rows(values_only=True), 1):
        if row_number == 1:
            continue
        if row and row[0] is not None:
            parsed_date = parse_date(row[0])
            if parsed_date:
                current_date = parsed_date
            elif isinstance(row[0], str) and row[0].strip() not in ("`", ""):
                current_category = row[0].strip()
        if not current_date or len(row) < 5:
            continue
        supplier, description = row[2], row[3]
        source_column_d_date = parse_date(description)
        total_with_vat = row[4] if len(row) > 4 and isinstance(row[4], (int, float)) else None
        amount = row[5] if len(row) > 5 and isinstance(row[5], (int, float)) else None
        vat_amount = row[6] if len(row) > 6 and isinstance(row[6], (int, float)) else None
        if total_with_vat is not None and vat_amount is not None:
            gross_amount = total_with_vat
        elif amount is not None:
            gross_amount = amount
        else:
            gross_amount = total_with_vat
        if not (supplier or description) or not isinstance(gross_amount, (int, float)):
            continue
        notes = "Imported from ACI CASH DISBURSEMENT 2026.xlsx"
        if source_column_d_date:
            notes += f"; source column D date: {source_column_d_date}"
        values = ["expense", current_date, row[1], None, str(supplier or description).strip(), None if source_column_d_date else description, expense_type(supplier, None if source_column_d_date else description, current_category), gross_amount, vat_amount or 0, row[8] if isinstance(row[8], (int, float)) else 0, row[10], row[11], expense_file.name, sheet.title, row_number, BATCH, notes]
        expense_row = f"({', '.join(sql(value) for value in values)})"
        rows.append(expense_row)
        expense_rows.append(expense_row)
        count += 1

lines = [
    "-- Generated from ACI legacy workbooks.",
    "-- Run with the Supabase service role in local development only.",
    f"INSERT INTO public.legacy_accounting_entries ({columns}) VALUES",
    ",\n".join(rows) + "\nON CONFLICT (source_workbook, source_sheet, source_row, import_batch) DO UPDATE SET entry_date = excluded.entry_date, reference_no = excluded.reference_no, document_no = excluded.document_no, counterparty_name = excluded.counterparty_name, description = excluded.description, expense_type = excluded.expense_type, gross_amount = excluded.gross_amount, vat_amount = excluded.vat_amount, ewt_amount = excluded.ewt_amount, tin = excluded.tin, address = excluded.address, notes = excluded.notes;",
    f"-- Imported rows: {count}",
]
out = ROOT / "supabase" / "snippets" / "legacy_accounting_import_local.sql"
out.parent.mkdir(parents=True, exist_ok=True)
out.write_text("\n".join(lines) + "\n", encoding="utf-8")
print(f"Generated {count} import rows: {out}")

expense_out = ROOT / "supabase" / "snippets" / "legacy_cash_disbursement_import.sql"
expense_lines = [
    "-- Generated only from ACI CASH DISBURSEMENT 2026.xlsx.",
    "-- Run the required legacy_accounting_entries migration before this import.",
    f"INSERT INTO public.legacy_accounting_entries ({columns}) VALUES",
    ",\n".join(expense_rows) + "\nON CONFLICT (source_workbook, source_sheet, source_row, import_batch) DO UPDATE SET entry_date = excluded.entry_date, reference_no = excluded.reference_no, document_no = excluded.document_no, counterparty_name = excluded.counterparty_name, description = excluded.description, expense_type = excluded.expense_type, gross_amount = excluded.gross_amount, vat_amount = excluded.vat_amount, ewt_amount = excluded.ewt_amount, tin = excluded.tin, address = excluded.address, notes = excluded.notes;",
    f"-- Imported rows: {len(expense_rows)}",
]
expense_out.write_text("\n".join(expense_lines) + "\n", encoding="utf-8")
print(f"Generated {len(expense_rows)} disbursement rows: {expense_out}")

pnl_file = Path(r"C:\Users\skhart\Downloads\ACI PROFIT AND LOSS 2026.xlsx")
pnl_workbook = openpyxl.load_workbook(pnl_file, data_only=True, read_only=True)
pnl_sheet = pnl_workbook.active
pnl_rows = []
month_columns = [(5, "2026-01-01"), (6, "2026-02-01"), (7, "2026-03-01")]
for row_number, row in enumerate(pnl_sheet.iter_rows(values_only=True), 1):
    category = row[0] if len(row) > 0 and row[0] else row[1] if len(row) > 1 else None
    detail = row[1] if len(row) > 1 and row[0] else None
    if not category:
        continue
    for column, period_start in month_columns:
        amount = row[column] if len(row) > column and isinstance(row[column], (int, float)) else None
        if amount is None:
            continue
        pnl_rows.append(
            "(" + ", ".join(sql(value) for value in [
                period_start,
                str(category).strip(),
                str(detail).strip() if detail else None,
                amount,
                pnl_file.name,
                pnl_sheet.title,
                row_number,
                BATCH,
            ]) + ")"
        )

pnl_out = ROOT / "supabase" / "snippets" / "legacy_profit_loss_lines_import_local.sql"
pnl_out.write_text(
    "insert into public.legacy_profit_loss_lines (period_start, category, detail, amount, source_workbook, source_sheet, source_row, import_batch) values\n"
    + ",\n".join(pnl_rows)
    + "\non conflict (source_workbook, source_sheet, source_row, period_start, import_batch) do update set category = excluded.category, detail = excluded.detail, amount = excluded.amount;\n",
    encoding="utf-8",
)
print(f"Generated {len(pnl_rows)} P&L detail rows: {pnl_out}")
