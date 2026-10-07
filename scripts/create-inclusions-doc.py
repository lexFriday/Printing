from docx import Document
from docx.shared import Inches, Pt, RGBColor
from pathlib import Path

doc = Document()
sec = doc.sections[0]
sec.top_margin = sec.bottom_margin = sec.left_margin = sec.right_margin = Inches(1)
doc.styles['Normal'].font.name = 'Calibri'
doc.styles['Normal'].font.size = Pt(11)
for name, size in [('Heading 1', 16), ('Heading 2', 13)]:
    style = doc.styles[name]
    style.font.name = 'Calibri'; style.font.size = Pt(size); style.font.color.rgb = RGBColor(20, 86, 160); style.font.bold = True

def para(text='', style=None, size=11, color=None, bold=False, italic=False):
    p = doc.add_paragraph(style=style)
    r = p.add_run(text); r.font.size = Pt(size); r.bold = bold; r.italic = italic
    if color: r.font.color.rgb = RGBColor(*color)
    p.paragraph_format.space_after = Pt(7); p.paragraph_format.line_spacing = 1.15
    return p
def bullets(items):
    for item in items: para(item, 'List Bullet')
def steps(items):
    for item in items: para(item, 'List Number')

para('What’s Included in the Printing System Today', size=17, color=(20,86,160), bold=True)
para('AutoMate PH · Current local build · 3 September 2026', size=10, color=(91,101,115), italic=True)
para('In one sentence', 'Heading 1')
para('The printing system keeps the workday in one flow: capture a customer request, prepare a quote, turn approved work into a Job Order, track money in and money out, invoice the customer, and review profitability.')
para('Main screens', 'Heading 1')

content = {
'Dashboard':'The Dashboard is the starting point for work that needs attention. It gives a quick view of operational and financial queues, such as quotes requiring action, unpaid invoices, pending expenses, and other high-level totals. What a user can see depends on their role.',
'Pending Expenses':'Pending Expenses is the main money-out entry point. A receipt can include a supplier/payee, description, category, payment source, petty-cash batch, Job Order, tax details, document number, amount, and payment status. Expenses may be general, linked to one Job Order, or shared across multiple Job Orders through amount-based allocations. The screen also supports spreadsheet import and expense editing.',
'Petty cash':'Petty cash works like a simple transaction register: create a cash-advance batch; add multiple receipts; link receipts to Job Orders; see starting cash, receipts, available/running balance, receipt count, and project totals; import receipts; close the batch; create one CV for the original cash received amount; and show Cash to return or Overage for review.',
'Customers':'Customers stores reusable customer information: customer name, complete address, TIN, notes, contact details, multiple contacts, search/list/detail/edit views, and related quote and invoice history.',
'Suppliers':'Suppliers is under Administration. It provides a reusable supplier master list with supplier name, TIN, address, contact person, contact number, email, notes, and active/inactive status. Suppliers can be selected during expense entry; historical expense snapshots remain tied to the original transaction details.',
'Item Templates':'Item Templates stores reusable quotation items and pricing defaults so common work can be added consistently to new quotes.',
'Users':'Users lets authorized administrators review accounts, assign roles, activate or deactivate access, and maintain user signatures used by printed documents where configured.',
'Audit Logs / Change History':'Audit Logs records who changed what and when. It covers important customer, quote, Job Order, invoice, expense, petty-cash, allocation, supplier, user, and workflow actions. Corrections can include the reason supplied by the user.'}

for heading, text in content.items():
    para(heading, 'Heading 2'); para(text)
    if heading == 'Pending Expenses':
        bullets(['General expenses with no Job Order.','One Job Order expenses.','Shared receipts allocated across multiple Job Orders.','Spreadsheet import, category filters, payment source, petty-cash linking, and editing.'])
    if heading == 'Petty cash':
        bullets(['Funding source and reference, custodian, and funding date.','Running balance and Job Order/project summary.','Closed-batch protection and one-CV-per-batch settlement.'])

para('Quotes', 'Heading 2'); para('Create, search, filter, edit, approve, correct, recover, and print customer quotations. Quotes support customer/contact details, dates, VAT/tax mode, terms, remarks, reference images, multiple line items, fixed-price or square-foot pricing, optional depth, and No charge items with a reason.')
para('Job Orders', 'Heading 2'); para('Convert approved quotes into Job Orders, view source details, track production through pending, in production, ready for delivery, delivered, and completed, correct details with permission, and print the Job Order and delivery receipt. Money-out entry is handled in Reports → Pending Expenses.')
para('Invoices', 'Heading 2'); para('Create invoices from completed Job Orders or directly, select billable items, invoice and due dates, tax information, record payment evidence, correct payment details with permission, and print invoices.')
para('Reports', 'Heading 2'); para('Reports includes Profit / P&L, Operations, Money In, Money Out, Pending Expenses, Disbursements, Petty cash, Cash sales, filters, and exports.')

para('Printing and exports', 'Heading 1'); para('Print views are available for quotations, Job Orders, delivery receipts, customer invoices, checks, and check vouchers. Long item descriptions, payees, addresses, project names, and references are fitted to the available print field. Calibration settings and offsets can be adjusted for the physical printer and paper sample. Spreadsheet exports include expense, supplier, category, Job Order/project, tax, document, allocation, amount, and total information where supported.')
para('Built-in safeguards', 'Heading 1'); bullets(['Role-based access and tenant-aware data access.','Required-field and amount validation.','No-charge quotation reasons.','Existing Job Order selection only.','Shared-expense allocation checks.','Petty-cash closure and one-CV protection.','Audit history for corrections and workflow actions.','Reports derived from operational records.','Draft recovery and preservation of values when a save fails.'])
para('Typical day-to-day flow', 'Heading 1'); steps(['Add or select the customer and contact.','Prepare and send a quote for approval.','Convert approved work into a Job Order and track production.','Enter money-out transactions in Reports → Pending Expenses.','Tag or allocate expenses to Job Orders.','Track petty-cash running balances.','Complete the Job Order and create the invoice.','Record payment and review disbursements, P&L, and exports.','Use Change History for corrections.'])
para('What this system does not replace', 'Heading 1'); para('The application is an operations and printing system. It does not replace a full external accounting platform such as QuickBooks. Tax, payment, approval, printer, and permission behavior depends on the configured workspace rules and the user’s role.')
para('Scope note', 'Heading 1'); para('This guide describes the screens and capabilities visible in the current local application build. Some actions are hidden from roles without permission, and Live may require separate deployment or schema configuration before it shows the same capabilities.')

out = Path('docs/current-printing-system-inclusions.docx'); out.parent.mkdir(exist_ok=True); doc.save(out)
